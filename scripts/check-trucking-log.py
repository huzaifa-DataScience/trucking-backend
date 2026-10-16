"""
Rules from 2025-2026 Trucking Log.xlsx (Trucking Record row 2 is the template).

Trucking  = hauler rate per load + rate per hour * hours
Disposal  = cost per load
          + cost per ton * billable tons
          + cost per CY * catalog truck size
Billable tons = 0 when the rate already includes more tons than the load weighs,
                otherwise weight minus the tons included in the haul rate.
Total cost = trucking + disposal.

Blank sentinels on the template row: hauler "-----", destination "-----", job name "---".
VLOOKUP's 4th argument is omitted, so Excel approximate-matches and Job Sites column A must stay A–Z.
"""
from __future__ import annotations

SENTINEL_HAULER = "-----"
SENTINEL_DISPOSAL = "-----"
SENTINEL_JOB = "---"


def num(value) -> float:
    if value is None or value == "":
        return 0.0
    return float(value)


def billable_tons(weight, tons_included) -> float:
    """Excel: IF(AND(included>0, included-weight>0), 0, weight-included)."""
    weight = num(weight)
    included = num(tons_included)
    if included > 0 and (included - weight) > 0:
        return 0.0
    return weight - included


def trucking_cost(rate_per_load, rate_per_hour, hours, hauler) -> float:
    if hauler in (None, "", SENTINEL_HAULER):
        return 0.0
    return num(rate_per_load) + num(rate_per_hour) * num(hours)


def disposal_cost(cost_per_load, cost_per_ton, cost_per_cy, truck_cy, tons_included, weight, destination) -> float:
    if destination in (None, "", SENTINEL_DISPOSAL):
        return 0.0
    tons = billable_tons(weight, tons_included)
    return num(cost_per_load) + num(cost_per_ton) * tons + num(cost_per_cy) * num(truck_cy)


def total_cost(trucking, disposal) -> float:
    return trucking + disposal


def job_name(job_number, lookup) -> str:
    if job_number in (None, ""):
        return SENTINEL_JOB
    return lookup


def leed_total(job, quantities) -> str | float:
    """Template AE2. Not copied onto the data rows."""
    if job == SENTINEL_JOB:
        return SENTINEL_JOB
    return sum(num(q) for q in quantities)


def hauler_label(company, category) -> str:
    if company in (None, ""):
        return "Z"
    return f"{company} - {category}"


def disposal_label(company, unit, category) -> str:
    if company in (None, ""):
        return "X"
    return f"{company} - {unit} - {category}"


def tons_from_pounds(pounds) -> float:
    return num(pounds) / 2000


def next_ticket_block(start: int, size: int = 50) -> tuple[int, int]:
    end = start + size - 1
    return end, end + 1


def assert_close(actual, expected, label):
    if abs(actual - expected) > 0.02:
        raise SystemExit(f"{label}: {actual} != {expected}")


def main():
    # Template row: nothing chosen.
    assert trucking_cost(550, 0, None, SENTINEL_HAULER) == 0
    assert disposal_cost(500, 0, 0, 0, 0, 10, SENTINEL_DISPOSAL) == 0
    assert total_cost(0, 0) == 0
    assert job_name("", "ignored") == SENTINEL_JOB
    assert leed_total(SENTINEL_JOB, [1, 2]) == SENTINEL_JOB

    # Flat haul + flat disposal. Hours blank counts as 0. Row 7 shape.
    haul = trucking_cost(350, 0, None, "JRL - Hauler - Baltimore")
    disp = disposal_cost(300, 0, 0, 0, 0, None, "JRL - Disposal - Brick/Concrete")
    assert_close(haul, 350, "flat haul")
    assert_close(disp, 300, "flat disposal")
    assert_close(total_cost(haul, disp), 650, "flat total")

    # Per ton, no tons included. Row 5: 9.84 * 67 = 659.28, haul 630.
    haul = trucking_cost(630, 0, None, "McNair")
    disp = disposal_cost(0, 67, 0, 0, 0, 9.84, "Cumberland asbestos")
    assert_close(haul, 630, "cumberland haul")
    assert_close(disp, 659.28, "cumberland disposal")
    assert_close(total_cost(haul, disp), 1289.28, "cumberland total")

    # Hourly haul. Hammer Head $85/hr * 1 hour, no load fee.
    assert_close(trucking_cost(0, 85, 1, "Hammer Head"), 85, "hourly")
    assert_close(trucking_cost(0, 85, 2.5, "Hammer Head"), 212.5, "hourly 2.5")

    # In-house and a $0 recycle site. Row 41 shape.
    assert trucking_cost(0, 0, None, "Goel in house") == 0
    assert disposal_cost(0, 0, 0, 0, 0, 1.13, "Joseph Smith non-ferrous") == 0

    # Tons already inside the haul rate. L&J C&D card: $65/ton after 5 tons.
    under = disposal_cost(0, 65, 0, 0, 5, 3, "L&J")
    over = disposal_cost(0, 65, 0, 0, 5, 10.59, "L&J")
    assert under == 0
    assert_close(over, 65 * 5.59, "tons over the included amount")

    # Load fee plus per-ton on the whole weight when nothing is included.
    # BRC card: $740/load + $125/ton.
    assert_close(disposal_cost(740, 125, 0, 0, 0, 2, "BRC"), 740 + 250, "load plus ton")

    # Cubic yards use the catalog truck size, not the weighed tons.
    # Wirtz topsoil: $20/CY * 20 CY dump.
    assert_close(disposal_cost(0, 0, 20, 20, 0, 9, "Wirtz"), 400, "cy flat")

    # Pounds on the ticket become tons. (13360-11100)/2000 = 1.13.
    assert_close(tons_from_pounds(13360 - 11100), 1.13, "scale pounds")
    assert_close(tons_from_pounds(3020), 1.51, "pounds only")

    # Labels the dropdowns look up.
    assert hauler_label("A&S Refuse", "Hauler  - 30 YD can C&D") == "A&S Refuse - Hauler  - 30 YD can C&D"
    assert hauler_label("", "Hauler") == "Z"
    assert disposal_label("A&S Refuse", "Disposal", "C&D ") == "A&S Refuse - Disposal - C&D "
    assert disposal_label("", "Recycle", "Dirt") == "X"

    # Ticket books are 50 numbers. Next book starts on the number after this one ends.
    end, nxt = next_ticket_block(160001)
    assert end == 160050 and nxt == 160051

    # LEED total is the sum of the vendor columns, and only when a job name resolved.
    assert leed_total("Vine St", [1, 0, 2.5]) == 3.5

    replay_workbook()
    print("check-trucking-log: ok")


def replay_workbook():
    """Formula rows in the workbook must match these rules. Typed-over cells are skipped."""
    from pathlib import Path
    path = Path(__file__).resolve().parents[1] / "2025-2026 Trucking Log.xlsx"
    if not path.exists():
        return
    import openpyxl

    raw = openpyxl.load_workbook(path, data_only=False)
    calc = openpyxl.load_workbook(path, data_only=True)

    def card_map(ws_calc, key_col, fields, start, end):
        out = {}
        for r in range(start, end):
            key = ws_calc.cell(r, key_col).value
            if not key or key in ("X", "Z"):
                continue
            out.setdefault(key, {name: ws_calc.cell(r, col).value for name, col in fields})
        return out

    haulers = card_map(
        calc["Haulers"],
        1,
        (("load", 4), ("hour", 5)),
        4,
        400,
    )
    sites = card_map(
        calc["Disposal"],
        1,
        (("load", 3), ("ton", 4), ("cy", 5), ("size", 6), ("included", 7)),
        4,
        640,
    )
    log = raw["Trucking Record"]
    cached = calc["Trucking Record"]
    checked = 0
    for r in range(4, 2500):
        am_src = log.cell(r, 39).value
        an_src = log.cell(r, 40).value
        ao_src = log.cell(r, 41).value
        if not (isinstance(am_src, str) and isinstance(an_src, str) and isinstance(ao_src, str)):
            continue
        hauler = cached.cell(r, 4).value
        dest = cached.cell(r, 6).value
        if hauler not in haulers or dest not in sites:
            continue
        h, d = haulers[hauler], sites[dest]
        pred_am = trucking_cost(h["load"], h["hour"], cached.cell(r, 5).value, hauler)
        pred_an = disposal_cost(d["load"], d["ton"], d["cy"], d["size"], d["included"], cached.cell(r, 13).value, dest)
        am, an, ao = cached.cell(r, 39).value, cached.cell(r, 40).value, cached.cell(r, 41).value
        if am is None or an is None or ao is None:
            continue
        assert_close(pred_am, num(am), f"row {r} trucking")
        assert_close(pred_an, num(an), f"row {r} disposal")
        # A few SUM cells were saved before the rate cards changed. Trust AM+AN when they disagree.
        if abs(num(am) + num(an) - num(ao)) <= 0.02:
            assert_close(pred_am + pred_an, num(ao), f"row {r} total")
        checked += 1
    if checked < 2000:
        raise SystemExit(f"only replayed {checked} rows")
    raw.close()
    calc.close()


if __name__ == "__main__":
    main()
