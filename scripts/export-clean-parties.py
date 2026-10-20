"""Turn Followup CRM invite lists into one contact per person.

The workbook joins several companies on one project with a comma and no space
("Inc.,Welch"). A comma followed by a space stays inside the name (", Inc.").
"""
import json
import re
from collections import Counter, defaultdict

import openpyxl

CRM = "Updated Data followup crm.xlsx"
COMPANIES = "companies-09042026.xlsx"
OUT = "/tmp/clean-parties.json"

LEGAL = re.compile(r"\b(incorporated|inc|llc|llp|corp|corporation|company|co|ltd)\b", re.I)
TYPE_ROLE = {
    "GC": "gc",
    "GenContMec": "gc",
    "Mechanical Cont": "mechanical_contractor",
    "Mec - SheetMetC": "mechanical_contractor",
    "Mec - Heat Trac": "mechanical_contractor",
    "Architect": "architect",
    "Building Owner": "owner",
    "Mechanical Engi": "mechanical",
    "Insulation Comp": "insulation",
}


def parts(value):
    s = "" if value is None else str(value).strip()
    if not s:
        return []
    return [p.strip(" ,") for p in re.split(r",(?=\S)", s) if p.strip(" ,")]


def norm(s):
    s = LEGAL.sub(" ", (s or "").lower().replace("&", " and "))
    return re.sub(r"[^a-z0-9]+", "", s)


def good_email(value):
    email = str(value or "").strip().lower()
    return email if re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email) else None


def good_contact(value):
    name = re.sub(r"\s+", " ", str(value or "")).strip(" -")
    if len(name) < 2 or "@" in name or "," in name:
        return None
    if re.match(r"^contact needed\b", name, re.I):
        return None
    return name[:200]


def good_phone(value):
    phone = re.sub(r"\s+", " ", str(value or "")).strip()
    digits = re.sub(r"\D", "", phone)
    if len(digits) < 10 or len(digits) > 15 or len(phone) > 40 or "," in phone:
        return None
    return phone


def good_company(value):
    company = re.sub(r"\s+", " ", str(value or "")).strip(" ,")
    if len(company) < 2 or "@" in company or re.search(r",(?=\S)", company):
        return None
    return company[:200]


def name_matches_email(contact, email):
    local = norm(email.split("@")[0])
    words = [norm(w) for w in re.sub(r"\(.*?\)", " ", contact).split()]
    words = [w for w in words if len(w) >= 3]
    if not words or not local:
        return False
    last, first = words[-1], words[0]
    return last in local or local in last or (first[0] + last) in local or (len(first) >= 4 and first in local)


def domain_key(email):
    return re.sub(r"[^a-z0-9]+", "", email.split("@")[-1].lower().split(".")[0])


def main():
    by_norm = {}
    by_email_master = {}
    wb = openpyxl.load_workbook(COMPANIES, read_only=True, data_only=True)
    for row in wb["Worksheet"].iter_rows(min_row=2, values_only=True):
        company = good_company(row[1])
        if not company:
            continue
        email = good_email(row[4])
        rec = {
            "company": company,
            "phone": good_phone(row[2]),
            "contact": good_contact(row[3]),
            "email": email,
            "role": TYPE_ROLE.get(str(row[11] or "").strip()),
        }
        by_norm[norm(company)] = rec
        if email:
            by_email_master[email] = rec
    wb.close()

    contact_votes = defaultdict(Counter)
    aligned = defaultdict(list)
    wb = openpyxl.load_workbook(CRM, read_only=True, data_only=True)
    for row in wb["Worksheet"].iter_rows(min_row=2, values_only=True):
        companies, contacts, emails, cells = parts(row[23]), parts(row[22]), parts(row[25]), parts(row[26])
        n = len(emails)
        for i, raw in enumerate(emails):
            email = good_email(raw)
            if not email:
                continue
            contact = good_contact(contacts[i]) if len(contacts) == n else None
            company = good_company(companies[i]) if len(companies) == n else None
            phone = good_phone(cells[i]) if len(cells) == n else None
            if contact:
                contact_votes[email][contact] += 1
            aligned[email].append((contact, company, phone, [c for c in companies if good_company(c)]))
    wb.close()

    def pick_contact(email):
        votes = contact_votes[email]
        master = by_email_master.get(email)
        matched = [(c, n) for c, n in votes.items() if name_matches_email(c, email)]
        if matched:
            return max(matched, key=lambda item: item[1])[0]
        if master and master["contact"] and name_matches_email(master["contact"], email):
            return master["contact"]
        if not votes:
            return master["contact"] if master else None
        contact, count = votes.most_common(1)[0]
        total = sum(votes.values())
        if count >= 3 and count >= total * 0.75:
            return contact
        return master["contact"] if master else None

    people = {}
    for email in set(contact_votes) | set(aligned) | set(by_email_master):
        contact = pick_contact(email)
        master = by_email_master.get(email)
        domain = domain_key(email)
        domain_votes = Counter()
        index_votes = Counter()
        phones = Counter()
        for row_contact, row_company, row_phone, companies in aligned.get(email, []):
            if contact and row_contact and norm(row_contact) != norm(contact):
                continue
            if row_company:
                index_votes[row_company] += 1
            if row_phone:
                phones[row_phone] += 1
            if len(domain) >= 4:
                for company in companies:
                    key = norm(company)
                    if key and (domain in key or (len(key) >= 4 and key in domain)):
                        domain_votes[company] += 1
        if master:
            company = master["company"]
        elif domain_votes:
            company = domain_votes.most_common(1)[0][0]
        elif index_votes:
            top, count = index_votes.most_common(1)[0]
            company = top if count >= 2 or sum(index_votes.values()) == 1 else None
        else:
            company = None
        if company and norm(company) in by_norm:
            company = by_norm[norm(company)]["company"]
        phone = phones.most_common(1)[0][0] if phones else None
        if not phone and master:
            phone = master["phone"]
        role = by_norm.get(norm(company or ""), {}).get("role") if company else None
        if not company and not contact:
            continue
        people[email] = {
            "company": company,
            "contactName": contact,
            "email": email,
            "phone": phone,
            "role": role,
        }

    for rec in by_norm.values():
        if rec["email"]:
            continue
        people["noemail:" + norm(rec["company"])] = {
            "company": rec["company"],
            "contactName": rec["contact"],
            "email": None,
            "phone": rec["phone"],
            "role": rec["role"],
        }

    rows = []
    for rec in people.values():
        roles = ["invite_contact"]
        if rec["role"] and rec["role"] != "invite_contact":
            roles.append(rec["role"])
        for role in roles:
            rows.append({**rec, "role": role})
    with open(OUT, "w") as f:
        json.dump(rows, f)
    print(f"clean parties {len(rows)} people {len(people)} -> {OUT}")


if __name__ == "__main__":
    main()
