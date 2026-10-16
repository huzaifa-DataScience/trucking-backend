"""Dump 2025-2026 Trucking Log.xlsx to /tmp/trucking-log-import.json for the operations import."""
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl


def hauler_label(company, category):
    if not company:
        return "Z"
    return f"{company} - {category}"


def disposal_label(company, unit, category):
    if not company:
        return "X"
    return f"{company} - {unit} - {category}"

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "2025-2026 Trucking Log.xlsx"
OUT = Path("/tmp/trucking-log-import.json")


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    s = str(value).strip()
    return s or None


def dec(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def main():
    book = openpyxl.load_workbook(SRC, data_only=True)
    haulers = []
    ws = book["Haulers"]
    for r in range(4, 400):
        company = text(ws.cell(r, 2).value)
        if not company:
            continue
        category = text(ws.cell(r, 3).value)
        haulers.append({
            "sourceRow": r,
            "company": company,
            "category": category,
            "label": text(ws.cell(r, 1).value) or hauler_label(company, category or ""),
            "ratePerLoad": dec(ws.cell(r, 4).value) or 0,
            "ratePerHour": dec(ws.cell(r, 5).value) or 0,
            "unit": text(ws.cell(r, 6).value),
            "showUpHours": dec(ws.cell(r, 7).value),
            "address": text(ws.cell(r, 8).value),
            "phone": text(ws.cell(r, 9).value),
            "contact": text(ws.cell(r, 10).value),
            "email": text(ws.cell(r, 11).value),
            "coi": text(ws.cell(r, 12).value),
        })

    disposal = []
    ws = book["Disposal"]
    for r in range(4, 640):
        company = text(ws.cell(r, 2).value)
        if not company:
            continue
        unit = text(ws.cell(r, 9).value)
        category = text(ws.cell(r, 10).value)
        disposal.append({
            "sourceRow": r,
            "company": company,
            "unit": unit,
            "category": category,
            "label": text(ws.cell(r, 1).value) or disposal_label(company, unit or "", category or ""),
            "costPerLoad": dec(ws.cell(r, 3).value) or 0,
            "costPerTon": dec(ws.cell(r, 4).value) or 0,
            "costPerCy": dec(ws.cell(r, 5).value) or 0,
            "truckSizeCy": dec(ws.cell(r, 6).value) or 0,
            "tonsIncluded": dec(ws.cell(r, 7).value) or 0,
            "address": text(ws.cell(r, 11).value),
            "phone": text(ws.cell(r, 12).value),
            "contact": text(ws.cell(r, 13).value),
            "email": text(ws.cell(r, 14).value),
        })

    # Trucking Record: row 3 is the header, data starts at row 4.
    qty_cols = {
        "lightTubes": 15,
        "metalsDirect": 16,
        "metalsReclaimed": 17,
        "insulation": 18,
        "wood": 19,
        "paper": 20,
        "aggregates": 21,
        "plastics": 22,
        "drywall": 23,
        "boardFeetWood": 24,
        "dirt": 25,
        "lumber": 26,
        "brickDeconstructed": 27,
        "brickPallets": 28,
        "roofing": 29,
        "landfill": 30,
        "rc57": 33,
        "rc6": 34,
        "rc2": 35,
        "cr57": 36,
        "dirtTons": 37,
    }
    hauls = []
    ws = book["Trucking Record"]
    for r in range(4, 2500):
        ticket_date = iso(ws.cell(r, 1).value)
        job_number = text(ws.cell(r, 2).value)
        hauler = text(ws.cell(r, 4).value)
        destination = text(ws.cell(r, 6).value)
        weight = dec(ws.cell(r, 13).value)
        tickets = [text(ws.cell(r, c).value) for c in (8, 9, 10, 11, 12)]
        if not any([ticket_date, job_number, hauler, destination, weight, *tickets]):
            continue
        if hauler in ("-----", "Who Hauled Materials"):
            hauler = None
        if destination in ("-----",):
            destination = None
        row = {
            "sourceRow": r,
            "ticketDate": ticket_date,
            "jobNumber": job_number,
            "jobName": text(ws.cell(r, 3).value),
            "haulerLabel": hauler,
            "hours": dec(ws.cell(r, 5).value),
            "disposalLabel": destination,
            "goelTicketNumber": tickets[0],
            "haulingTicketNumber": tickets[1],
            "disposalTicketNumber": tickets[2],
            "wasteManifestNumber": tickets[3],
            "transferJobNumber": tickets[4],
            "weightTons": weight,
        }
        for name, col in qty_cols.items():
            row[name] = dec(ws.cell(r, col).value)
        if not row["ticketDate"]:
            continue
        hauls.append(row)

    OUT.write_text(json.dumps({"haulers": haulers, "disposal": disposal, "hauls": hauls}))
    print(f"exported haulers={len(haulers)} disposal={len(disposal)} hauls={len(hauls)}")


if __name__ == "__main__":
    main()
